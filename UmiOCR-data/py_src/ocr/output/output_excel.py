from .output import Output
from .tools import getDataText
import os
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill


class OutputExcel(Output):
    def __init__(self, argd):
        super().__init__(argd)
        self.dir = argd["outputDir"]  # 输出路径（文件夹）
        self.fileName = argd["outputFileName"]  # 文件名
        self.outputPath = f"{self.dir}/{self.fileName}.xlsx"  # 输出路径
        self.ignoreBlank = argd["ignoreBlank"]  # 忽略空白文件
        self.includeConfidence = argd.get("includeConfidence", False)  # 是否包含置信度

        # 初始化Excel工作簿和工作表
        self.initExcelFile()

    def initExcelFile(self):
        """初始化Excel文件"""
        # 创建工作簿
        self.wb = Workbook()

        # 移除默认的工作表
        default_sheet = self.wb.get_sheet_by_name("Sheet")
        self.wb.remove_sheet(default_sheet)

        # 创建新的工作表
        self.sheet = self.wb.create_sheet(title="OCR识别结果")

        # 设置表头
        self.set_header()

        # 初始化行计数器
        self.row_count = 2  # 第一行是表头

    def set_header(self):
        """设置Excel表格的表头"""
        # 定义表头列
        headers = ["图片路径", "识别文字"]

        # 如果需要包含置信度，则添加置信度列
        if self.includeConfidence:
            headers.append("置信度")

        # 写入表头
        for col, header in enumerate(headers, start=1):
            cell = self.sheet.cell(row=1, column=col)
            cell.value = header

            # 设置表头样式
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = Border(left=Side(style="thin"), right=Side(style="thin"), top=Side(style="thin"), bottom=Side(style="thin"))

        # 调整列宽
        self.sheet.column_dimensions["A"].width = 50
        self.sheet.column_dimensions["B"].width = 100
        if self.includeConfidence:
            self.sheet.column_dimensions["C"].width = 15

    def print(self, res):  # 输出图片信息
        if not res["code"] == 100 and self.ignoreBlank:
            return  # 忽略空白图片

        # 图片路径
        image_path = res["path"]

        # 识别文字
        if res["code"] == 100:
            text = getDataText(res["data"])
        elif res["code"] == 101:
            text = "无文字"
        else:
            text = f"错误原因：{res['data']}"

        # 置信度
        confidence = ""
        if self.includeConfidence and res["code"] == 100:
            # 这里可以根据实际的OCR结果数据结构来获取置信度
            # 假设res["data"]是一个包含置信度的列表
            if isinstance(res["data"], list) and len(res["data"]) > 0:
                # 取平均置信度
                avg_confidence = sum(item.get("confidence", 0) for item in res["data"]) / len(res["data"])
                confidence = f"{avg_confidence:.2f}%"

        # 写入数据行
        self.sheet.cell(row=self.row_count, column=1).value = image_path
        self.sheet.cell(row=self.row_count, column=2).value = text
        if self.includeConfidence:
            self.sheet.cell(row=self.row_count, column=3).value = confidence

        # 设置数据行样式
        for col in range(1, len(self.sheet.columns) + 1):
            cell = self.sheet.cell(row=self.row_count, column=col)
            cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            cell.border = Border(left=Side(style="thin"), right=Side(style="thin"), top=Side(style="thin"), bottom=Side(style="thin"))

        # 增加行计数器
        self.row_count += 1

    def onEnd(self):  # 结束输出。
        # 确保输出目录存在
        os.makedirs(self.dir, exist_ok=True)

        # 保存Excel文件
        self.wb.save(self.outputPath)

